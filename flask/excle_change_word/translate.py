from docx import Document
import xlrd

document = Document()


def create_word_table(num, filename):
    for i in range(1,num):
        table = document.add_table(rows=20, cols=4, style='Light Grid')
        table.cell(1, 1).merge(table.cell(1, 3))
        table.cell(2, 1).merge(table.cell(2, 3))
        table.cell(3, 0).merge(table.cell(5, 0))
        table.cell(3, 2).merge(table.cell(3, 3))
        table.cell(4, 2).merge(table.cell(4, 3))
        table.cell(5, 2).merge(table.cell(5, 3))
        table.cell(6, 1).merge(table.cell(6, 3))
        table.cell(7, 1).merge(table.cell(7, 3))
        table.cell(8, 0).merge(table.cell(8, 3))
        table.cell(16, 1).merge(table.cell(16, 3))
        table.cell(17, 1).merge(table.cell(17, 3))
        table.cell(18, 1).merge(table.cell(18, 3))
        w_cells = table.rows[0].cells
        w_cells[0].text = u'用例名称'
        w_cells[2].text = u'标识'
        w_cells = table.rows[1].cells
        w_cells[0].text = u'测试追踪'
        w_cells[1].text = u'无'
        w_cells = table.rows[2].cells
        w_cells[0].text = u'用例综述'
        w_cells = table.rows[3].cells
        w_cells[1].text = u'硬件配置'
        w_cells[2].text = u'无'
        w_cells = table.rows[4].cells
        w_cells[0].text = u'用例的初始化'
        w_cells[1].text = u'软件配置'
        w_cells[2].text = u'无'
        w_cells = table.rows[5].cells
        w_cells[1].text = u'测试配置'
        w_cells[2].text = u'无'
        w_cells = table.rows[6].cells
        w_cells[0].text = u'前提和约束'
        w_cells = table.rows[7].cells
        w_cells[0].text = u'设计方法'
        w_cells = table.rows[8].cells
        w_cells[0].text = u'执行步骤'
        w_cells = table.rows[9].cells
        w_cells[0].text = u'序号'
        w_cells[1].text = u'输入及操作'
        w_cells[2].text = u'期望结果与评价标指'
        w_cells[3].text = u'实测结果'
        for j in range(1, 7):
            w_cells = table.rows[9 + j].cells
            w_cells[0].text = u'%d' % j
        w_cells = table.rows[16].cells
        w_cells[0].text = u'过程终止条件'
        w_cells[1].text = u'所有测试步骤执行完毕或因某种原因导致测试步骤无法执行'
        w_cells = table.rows[17].cells
        w_cells[0].text = u'结果评估'
        w_cells[1].text = u'该用例测试通过。'
        w_cells = table.rows[18].cells
        w_cells[0].text = u'备注'
        w_cells[1].text = u'无'
        w_cells = table.rows[19].cells
        w_cells[0].text = u'测试人员'
        # w_cells[1].text = u''
        w_cells[2].text = u'测试时间'
        w_cells[3].text = u'20201225'
        sheets = read_excel()
        sheet = sheets[6]
        table.cell(0, 1).text = sheet.cell_value(i, 3)  # 用例名称
        table.cell(0, 3).text = sheet.cell_value(i, 0)  # 标识
        table.cell(2, 1).text = sheet.cell_value(i, 2)  # 用例综述
        table.cell(6, 1).text = sheet.cell_value(i, 4)  # 前提与约束
        step_list = sheet.cell_value(i, 5).split('\n')
        step_list_len = len(step_list)
        for k in range(step_list_len):
            st = step_list[k].strip()[2:]
            table.cell(10 + k, 1).text = st  # 测试步骤
            table.cell(10 + k, 3).text = u'通过'  # 测试结果
        resp_list = sheet.cell_value(i, 6).split('\n')  # 期望步骤
        if len(resp_list) > 1:
            for m in range(len(resp_list)):
                resp = resp_list[m].strip()[2:]
                table.cell(10 + m, 2).text = resp
        else:
            table.cell(10, 2).merge(table.cell(9 + step_list_len, 2))
            table.cell(10, 2).text = resp_list[0].strip()
        document.add_page_break()
    document.save('%s.docx' % filename)


def read_excel():
    wb = xlrd.open_workbook(r'系统监控软件-测试用例v1.0.xlsx')
    sheetnames = wb.sheet_names()[1:]
    sheets = []
    for sheetname in sheetnames:
        sheet = wb.sheet_by_name(sheetname)
        sheets.append(sheet)
    return sheets


def write_word():
    sheets = read_excel()
    sheet = sheets[0]
    return sheet


if __name__ == '__main__':
    # create_word_table(47, '性能指标评估')
    # create_word_table(13, '作战使用准则合成与分析')
    # create_word_table(13, '作战可行性分析评估')
    # create_word_table(15, '辅助功能')
    create_word_table(3, '状态数据上报')
    # sheets = read_excel()
    # for i in sheets:
    #     print(i.name)
